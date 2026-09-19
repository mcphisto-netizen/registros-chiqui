#!/usr/bin/env python3
"""Regenera excel/ a partir de los CSVs de datos/."""
import csv
from collections import defaultdict
from datetime import date
from statistics import mean, pstdev
from scipy import stats
from openpyxl import Workbook

def read(path):
    with open(path, newline="") as f:
        return [{k: r[k] for k in r} for r in csv.DictReader(f)]

def sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    for c, h in enumerate(rows[0], 1):
        ws.cell(1, c, h)
    for i, r in enumerate(rows[1:], 2):
        for c, v in enumerate(r, 1):
            ws.cell(i, c, v)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 13
    return ws

# ---------- 01 glucosa ----------
g = read("datos/glucosa.csv")
gs = [["fecha", "momento", "valor", "origen"]]
for r in g:
    gs.append([r["fecha"], r["momento"], float(r["valor"].replace(",", ".")), r.get("origen", "")])

# ---------- 02 medicacion ----------
m = read("datos/medicacion.csv")
ms = [["fecha", "momento", "farmaco", "dosis", "estado", "origen"]]
for r in m:
    ms.append([r["fecha"], r["momento"], r["farmaco"], r["dosis"], r.get("estado", ""), r.get("origen", "")])

# ---------- 03 comidas ----------
c = read("datos/comidas.csv")
cs = [["fecha", "momento", "ingrediente", "dosis", "estado", "origen"]]
for r in c:
    dos = r["dosis"]
    try: dos = float(dos.replace(",", "."))
    except ValueError: pass
    cs.append([r["fecha"], r["momento"], r["ingrediente"], dos, r.get("estado", ""), r.get("origen", "")])

# ---------- 04 analisis ----------
glu = []
for r in g:
    v = r["valor"].replace(",", ".")
    if not v.replace(".", "", 1).isdigit(): continue
    y, mo, d = map(int, r["fecha"].split("-"))
    glu.append((date(y, mo, d), r["momento"], float(v)))
glu.sort()

metricas = [["Mes", "n", "TIR% 100-250", "TAR% >250", "TBR% <80", "TBR% <70", "CV%", "media", "SD"]]
for mo in range(3, 10):
    v = [x[2] for x in glu if x[0].month == mo]
    if not v: continue
    metricas.append([mo, len(v),
        sum(100 <= x <= 250 for x in v) / len(v) * 100,
        sum(x > 250 for x in v) / len(v) * 100,
        sum(x < 80 for x in v) / len(v) * 100,
        sum(x < 70 for x in v) / len(v) * 100,
        pstdev(v) / mean(v) * 100, mean(v), pstdev(v)])

dias = defaultdict(dict)
for dt, mom, v in glu:
    if mom in ("mediodia", "medianoche"): dias[dt][mom] = v
pairs = [(dias[d]["mediodia"], dias[d]["medianoche"]) for d in dias if "mediodia" in dias[d] and "medianoche" in dias[d]]
diff = [p[1] - p[0] for p in pairs]
w = stats.wilcoxon(diff)
brecha = [["Parametro", "Valor"],
    ["pares (días con ambos horarios)", len(pairs)],
    ["media dif (medianoche - mediodía) mg/dL", mean(diff)],
    ["SD de la dif", pstdev(diff)],
    ["Wilcoxon p", w.pvalue],
    ["% días noche > mediodía", sum(d > 0 for d in diff) / len(diff) * 100]]

lows = sorted([x for x in glu if x[2] < 80])
glu_s = sorted(glu)
reb = [["Fecha", "Momento", "Valor<80", "Siguiente lectura", "Delta", "Siguiente momento"]]
for i, x in enumerate(glu_s[:-1]):
    if x[2] < 80:
        nxt = glu_s[i + 1]
        reb.append([x[0].isoformat(), x[1], x[2], nxt[2], nxt[2] - x[2], nxt[1]])

a = [x[2] for x in glu]
recent = [x[2] for x in glu if x[0] >= date(2026, 8, 24)]
fru = [["Parametro", "Valor"],
    ["Ecuación (Kang 2015)", "eAG (mg/dL) = 0,59 x Fructosamina (µmol/L) - 9,6  =>  F = (eAG + 9,6) / 0,59"],
    ["Media glucosa período completo", mean(a)],
    ["F estimada período completo (µmol/L)", (mean(a) + 9.6) / 0.59],
    ["Media últimas 4 sem (24/8-19/9)", mean(recent)],
    ["F estimada últimas 4 sem (µmol/L)", (mean(recent) + 9.6) / 0.59]]

lab = [["fecha", "parametro", "valor", "unidad", "ref_min", "ref_max", "notas"]]
for r in read("datos/laboratorio.csv"):
    val = r.get("valor", "")
    try: val = float(val.replace(",", "."))
    except ValueError: pass
    lab.append([r.get("fecha", ""), r.get("parametro", ""), val, r.get("unidad", ""),
                r.get("ref_min", ""), r.get("ref_max", ""), r.get("notas", "")])

biblio = [["Tema", "Referencia / DOI"],
    ["Guías AAHA perros/cats", "AAHA 2018 Diabetes Management Guidelines (PMID 29314873)"],
    ["Objetivo 100-250 mg/dL + Somogyi", "Merck/Vetsulin: glucose curves y Somogyi effect (merck-animal-health-usa.com)"],
    ["Glucosa nocturna > diurna en perros", "JVIM 2021, Assessment of postprandial hyperglycemia and circadian fluctuation (FGMS), DOI 10.1111/jvim.16060"],
    ["Métricas CGM en perros diabéticos", "FreeStyle Libre metrics in diabetic dogs, PMC12175195 (TIR/TAR/media/CV)"],
    ["eAG = 0,59x F - 9,6", "Kang DS 2015, Canine fructosamine, PMC4397269"],
    ["Cut-offs caninos de fructosamina", "GOOD <360 / FAIR 360-442 / POOR >=443, PMC8880912"],
    ["Fructosamina e hipoglucemias", "Kuzi S 2023, Vet Record, DOI 10.1002/vetr.2236"],
    ["CV <36% umbral riesgo de hipo", "Castañeda 2023 DOI 10.1111/dom.15139; Mo 2020 PMC8169344"],
    ["Variabilidad día-a-día de insulina", "JVIM 2021, day-to-day CV entre formulaciones, DOI 10.1111/jvim.16006"],
    ["Antibióticos en diarrea crónica canina", "Metronidazol/tilosina como elección; amoxicilina no primera línea, PMC7079140"],
    ["Cobalamina (B12) en enteropatía crónica", "Déficit = mala prognosis; suplementación mejora clínica, PMC11898182"]]

leyenda = [["Abreviatura", "Significado"],
    ["TIR", "Tiempo dentro del rango: % de lecturas en 100-250 mg/dL"],
    ["TAR", "Tiempo por encima del rango: % de lecturas > 250 mg/dL"],
    ["TBR", "Tiempo por debajo del rango: % de lecturas < 80 mg/dL"],
    ["CV", "Coeficiente de variación: desviación estándar / media x 100"],
    ["SD", "Desviación estándar"],
    ["GLU / F", "Glucosa (planillas) / Fructosamina (µmol/L)"],
    ["eAG", "Glucosa promedio estimada (estimated average glucose)"],
    ["CGM / BG", "Monitorización continua de glucosa / Glucosa sanguínea"],
    ["UCCR", "Cociente cortisol-creatinina urinaria (control del Cushing)"],
    ["UPC", "Cociente proteína-creatinina urinaria (proteinuria)"],
    ["LPCE (cPLI)", "Lipasa pancreática específica canina"],
    ["Hto / Hb", "Hematocrito / Hemoglobina"],
    ["T4 libre / TSH", "Tiroxina libre / Hormona estimulante de la tiroides"],
    ["ALT / FAL", "Enzimas hepáticas: alanina-aminotransferasa / fosfatasa alcalina"]]

notas = [["Aviso", "Estas hojas de análisis fueron generadas por inteligencia artificial a partir de las bases de datos domiciliarias y la bibliografía citada. No constituyen diagnóstico ni sustituyen el criterio veterinario. Los valores domiciliarios son autorreportados."]]

for fname, name, rows in [("excel/01_glucosa.xlsx", "Glucosa", gs),
                          ("excel/02_medicacion.xlsx", "Medicacion", ms),
                          ("excel/03_comidas.xlsx", "Comidas", cs)]:
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, name, rows)
    wb.save(fname)
    print("OK", fname)

wb = Workbook()
wb.remove(wb.active)
sheet(wb, "Metricas", metricas)
sheet(wb, "Brecha noche-dia", brecha)
sheet(wb, "Rebote post-hipo", reb)
sheet(wb, "Fructosamina est.", fru)
sheet(wb, "Laboratorio", lab)
sheet(wb, "Bibliografia", biblio)
sheet(wb, "Leyenda", leyenda)
sheet(wb, "Notas IA", notas)
wb.save("excel/04_analisis.xlsx")
print("OK excel/04_analisis.xlsx")